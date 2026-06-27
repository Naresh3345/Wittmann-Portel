import os
import csv
import re
import secrets
import smtplib
from email.message import EmailMessage
from functools import wraps
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

import psycopg
from dotenv import load_dotenv
from flask import Flask, flash, redirect, render_template, request, send_file, session, url_for
from psycopg.rows import dict_row
from psycopg.types.json import Jsonb
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename


load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
MAIN_PROJECT_DIR = Path(os.getenv("INTERVIEW_PROJECT_DIR", r"E:\wittmann_interview_ai"))
REPORT_DIR = Path(os.getenv("REPORT_DIR", str(MAIN_PROJECT_DIR / "reports"))).resolve()
RESUME_DIR = (REPORT_DIR / "resumes").resolve()

DEFAULT_DATABASE_URL = "postgresql://postgres:postgres@localhost:5432/wittmann_interview_ai"
DATABASE_URL = os.getenv("DATABASE_URL", DEFAULT_DATABASE_URL)
COMPANY_NAME = os.getenv("COMPANY_NAME", "WITTMANN BATTENFELD India Pvt. Ltd.")
HR_USERNAME = os.getenv("HR_USERNAME", "hr")
HR_PASSWORD_HASH = os.getenv("HR_PASSWORD_HASH", generate_password_hash(os.getenv("HR_PASSWORD", "Wittmann@123")))
SMTP_HOST = os.getenv("SMTP_HOST", "").strip()
SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))
SMTP_USER = os.getenv("SMTP_USER", "").strip()
SMTP_PASSWORD = os.getenv("SMTP_PASSWORD", "").replace(" ", "").strip()
SMTP_FROM = os.getenv("SMTP_FROM", SMTP_USER).strip()
HR_PORTAL_URL = os.getenv("HR_PORTAL_URL", "http://192.168.70.102:5050").strip()
ENV_PATH = BASE_DIR / ".env"

SECTION_COUNTS = {"Aptitude": 15, "Programming": 3}
SHORTLIST_MIN_SCORE = float(os.getenv("SHORTLIST_MIN_SCORE", "70"))
QUESTION_IMPORT_COLUMNS = {
    "role_name": "Role Name",
    "role_slug": "Role Slug",
    "question_code": "Question Code",
    "section": "Section",
    "topic": "Topic",
    "difficulty": "Difficulty",
    "question_text": "Question",
    "options": "Options",
    "correct_answer": "Correct Answer",
    "keywords": "Keywords",
    "marks": "Marks",
    "active": "Active",
}

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "change-this-hr-portal-secret")


def mask_dsn_password(dsn):
    parsed = urlsplit(dsn)
    if not parsed.password:
        return dsn
    netloc = f"{parsed.username}:***@{parsed.hostname}"
    if parsed.port:
        netloc += f":{parsed.port}"
    return urlunsplit((parsed.scheme, netloc, parsed.path, parsed.query, parsed.fragment))


def get_db():
    return psycopg.connect(DATABASE_URL, row_factory=dict_row, connect_timeout=5)


def update_env_file(updates):
    lines = []
    seen = set()
    if ENV_PATH.exists():
        lines = ENV_PATH.read_text(encoding="utf-8", errors="ignore").splitlines()
    updated_lines = []
    for line in lines:
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in line:
            updated_lines.append(line)
            continue
        key, _ = line.split("=", 1)
        key = key.strip()
        if key in updates:
            updated_lines.append(f"{key}={updates[key]}")
            seen.add(key)
        else:
            updated_lines.append(line)
    for key, value in updates.items():
        if key not in seen:
            updated_lines.append(f"{key}={value}")
    ENV_PATH.write_text("\n".join(updated_lines).rstrip() + "\n", encoding="utf-8")


def apply_smtp_settings(settings):
    global SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASSWORD, SMTP_FROM
    SMTP_HOST = settings["SMTP_HOST"].strip()
    SMTP_PORT = int(settings["SMTP_PORT"] or "587")
    SMTP_USER = settings["SMTP_USER"].strip()
    SMTP_PASSWORD = settings["SMTP_PASSWORD"].replace(" ", "").strip()
    SMTP_FROM = settings["SMTP_FROM"].strip() or SMTP_USER


@app.errorhandler(psycopg.OperationalError)
def database_connection_error(error):
    return render_template(
        "database_error.html",
        company_name=COMPANY_NAME,
        db_label=mask_dsn_password(DATABASE_URL),
        error=str(error),
    ), 503


def require_login(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("hr_logged_in"):
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)

    return wrapped


def require_admin(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("hr_logged_in"):
            return redirect(url_for("login", next=request.path))
        if not session.get("hr_is_admin"):
            flash("Only the main HR admin can manage portal users.", "error")
            return redirect(url_for("dashboard"))
        return view(*args, **kwargs)

    return wrapped


def ensure_hr_user_table(conn):
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS hr_portal_users (
            user_id SERIAL PRIMARY KEY,
            username TEXT NOT NULL UNIQUE,
            email TEXT NOT NULL UNIQUE,
            full_name TEXT NOT NULL,
            password_hash TEXT NOT NULL,
            is_admin BOOLEAN NOT NULL DEFAULT FALSE,
            active BOOLEAN NOT NULL DEFAULT TRUE,
            must_change_password BOOLEAN NOT NULL DEFAULT TRUE,
            created_at TIMESTAMPTZ NOT NULL DEFAULT now(),
            updated_at TIMESTAMPTZ NOT NULL DEFAULT now()
        )
        """
    )
    admin = conn.execute("SELECT 1 FROM hr_portal_users WHERE username = %s", (HR_USERNAME,)).fetchone()
    if not admin:
        default_email = os.getenv("HR_EMAIL", f"{HR_USERNAME}@localhost").strip()
        conn.execute(
            """
            INSERT INTO hr_portal_users
                (username, email, full_name, password_hash, is_admin, active, must_change_password)
            VALUES (%s, %s, %s, %s, TRUE, TRUE, FALSE)
            ON CONFLICT (username) DO NOTHING
            """,
            (HR_USERNAME, default_email, "Main HR Admin", HR_PASSWORD_HASH),
        )


def send_hr_access_email(user, password):
    if not all([SMTP_HOST, SMTP_USER, SMTP_PASSWORD, SMTP_FROM]):
        return {"sent": False, "status": "not_configured", "error": "SMTP settings are incomplete."}

    login_url = HR_PORTAL_URL or request.url_root.rstrip("/")
    message = EmailMessage()
    message["Subject"] = f"{COMPANY_NAME} HR Portal Access"
    message["From"] = SMTP_FROM
    message["To"] = user["email"]
    message.set_content(
        "\n".join(
            [
                f"Hello {user['full_name']},",
                "",
                f"You have been given access to the {COMPANY_NAME} HR Portal.",
                f"Login URL: {login_url}",
                f"Username: {user['username']}",
                f"Password: {password}",
                "",
                "Please sign in and keep this access secure.",
            ]
        )
    )

    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=15) as server:
        server.starttls()
        server.login(SMTP_USER, SMTP_PASSWORD)
        server.send_message(message)
    return {"sent": True, "status": "sent", "error": ""}


def send_candidate_shortlist_email(candidate):
    if not all([SMTP_HOST, SMTP_USER, SMTP_PASSWORD, SMTP_FROM]):
        return {"sent": False, "status": "not_configured", "error": "SMTP settings are incomplete."}

    message = EmailMessage()
    message["Subject"] = f"Shortlisted for Second Round Technical Interview - {COMPANY_NAME}"
    message["From"] = SMTP_FROM
    message["To"] = candidate["email"]
    message.set_content(
        "\n".join(
            [
                f"Dear {candidate['name']},",
                "",
                f"Greetings from {COMPANY_NAME}.",
                "",
                "We are pleased to inform you that you have been shortlisted based on your interview assessment.",
                f"You have been selected for the second round technical interview for the {candidate['role_name']} role.",
                "",
                "Our HR team will contact you shortly with the schedule and further details for the technical round.",
                "",
                "Please keep your phone and email available for communication.",
                "",
                "Congratulations, and we wish you the best for the next round.",
                "",
                "Regards,",
                "HR Team",
                COMPANY_NAME,
            ]
        )
    )

    with smtplib.SMTP(SMTP_HOST, SMTP_PORT, timeout=15) as server:
        server.starttls()
        server.login(SMTP_USER, SMTP_PASSWORD)
        server.send_message(message)
    return {"sent": True, "status": "sent", "error": ""}


def friendly_smtp_error(error):
    message = str(error)
    lowered = message.lower()
    if "username and password not accepted" in lowered or "badcredentials" in lowered:
        return "Gmail rejected the SMTP login. Use a Google App Password, not the normal Gmail password, and make sure SMTP User and From Email are the same Gmail account."
    if "smtp settings are incomplete" in lowered:
        return "SMTP settings are incomplete. Open SMTP Settings and fill host, port, user, app password, and from email."
    if "authentication" in lowered or "authenticate" in lowered:
        return "SMTP authentication failed. Check the email address and app password in SMTP Settings."
    return f"Email could not be sent. Check SMTP Settings. Details: {message}"


def safe_project_file(path, allowed_root):
    if not path:
        return None
    candidate = Path(path)
    if not candidate.is_absolute():
        candidate = (MAIN_PROJECT_DIR / candidate).resolve()
    else:
        candidate = candidate.resolve()
    try:
        candidate.relative_to(allowed_root)
    except ValueError:
        return None
    return candidate if candidate.exists() else None


def fetch_dashboard_stats():
    with get_db() as conn:
        stats = conn.execute(
            """
            SELECT
                COUNT(*) AS total_interviews,
                COUNT(*) FILTER (WHERE i.status = 'Completed') AS completed_interviews,
                COUNT(*) FILTER (WHERE i.shortlist_status = 'Shortlisted') AS shortlisted,
                ROUND(COALESCE(AVG(i.total_score), 0)::numeric, 2) AS average_score
            FROM interviews i
            """
        ).fetchone()
        role_rows = conn.execute(
            """
            SELECT r.role_name, COUNT(i.interview_id) AS total,
                   COUNT(*) FILTER (WHERE i.shortlist_status = 'Shortlisted') AS shortlisted
            FROM roles r
            LEFT JOIN interviews i ON i.role_id = r.role_id
            GROUP BY r.role_name
            ORDER BY r.role_name
            """
        ).fetchall()
        recent = fetch_reports(conn, limit=8)
    return stats, role_rows, recent


def fetch_reports(conn, status="", role_id="", search="", limit=None):
    filters = ["i.report_path IS NOT NULL"]
    params = []
    if status:
        filters.append("i.shortlist_status = %s")
        params.append(status)
    if role_id:
        filters.append("i.role_id = %s")
        params.append(int(role_id))
    if search:
        filters.append("(u.name ILIKE %s OR u.email ILIKE %s OR u.phone ILIKE %s)")
        like = f"%{search}%"
        params.extend([like, like, like])
    limit_sql = "LIMIT %s" if limit else ""
    if limit:
        params.append(limit)
    return conn.execute(
        f"""
        SELECT i.interview_id, i.date, i.total_score, i.status, i.shortlist_status,
               i.shortlist_reason, i.report_path,
               u.name, u.email, u.phone, u.resume_path, r.role_id, r.role_name
        FROM interviews i
        JOIN users u ON u.user_id = i.user_id
        JOIN roles r ON r.role_id = i.role_id
        WHERE {" AND ".join(filters)}
        ORDER BY i.date DESC
        {limit_sql}
        """,
        tuple(params),
    ).fetchall()


def group_reports_by_shortlist(reports, mode="automatic"):
    shortlisted = []
    not_shortlisted = []
    for report in reports:
        score = float(report.get("total_score") or 0)
        stored_status = report.get("shortlist_status") or "Pending"
        manually_shortlisted = stored_status == "Shortlisted"
        automatically_shortlisted = stored_status not in {"Rejected", "Needs Review"} and score >= SHORTLIST_MIN_SCORE
        is_shortlisted = automatically_shortlisted if mode == "automatic" else manually_shortlisted
        if is_shortlisted:
            report["display_shortlist_status"] = "Shortlisted"
            report["display_shortlist_reason"] = report.get("shortlist_reason") or (
                f"Candidate met the {SHORTLIST_MIN_SCORE:g} marks benchmark."
                if mode == "automatic"
                else "Candidate was manually approved by HR for the second round."
            )
            shortlisted.append(report)
        else:
            report["display_shortlist_status"] = stored_status if stored_status != "Shortlisted" else "Not Shortlisted"
            report["display_shortlist_reason"] = report.get("shortlist_reason") or (
                f"Candidate did not meet the {SHORTLIST_MIN_SCORE:g} marks benchmark."
                if mode == "automatic"
                else "Waiting for HR manual approval."
            )
            not_shortlisted.append(report)
    return shortlisted, not_shortlisted


def fetch_roles(conn):
    return conn.execute("SELECT role_id, role_slug, role_name FROM roles ORDER BY role_name").fetchall()


def fetch_role_question_counts(conn):
    return conn.execute(
        """
        SELECT
            r.role_id,
            r.role_slug,
            r.role_name,
            COUNT(q.question_id) AS total_questions,
            COUNT(q.question_id) FILTER (WHERE q.active = TRUE) AS active_questions,
            COUNT(q.question_id) FILTER (WHERE q.section = 'Aptitude' AND q.active = TRUE) AS active_aptitude,
            COUNT(q.question_id) FILTER (WHERE q.section = 'Programming' AND q.active = TRUE) AS active_programming
        FROM roles r
        LEFT JOIN question_bank q ON q.role_slug = r.role_slug
        GROUP BY r.role_id, r.role_slug, r.role_name
        ORDER BY r.role_name
        """
    ).fetchall()


def fetch_questions(conn, role_slug="", section="", active=""):
    filters = []
    params = []
    if role_slug:
        filters.append("role_slug = %s")
        params.append(role_slug)
    if section:
        filters.append("section = %s")
        params.append(section)
    if active in {"true", "false"}:
        filters.append("active = %s")
        params.append(active == "true")
    where_sql = f"WHERE {' AND '.join(filters)}" if filters else ""
    return conn.execute(
        f"""
        SELECT *
        FROM question_bank
        {where_sql}
        ORDER BY role_slug, CASE section WHEN 'Aptitude' THEN 1 WHEN 'Programming' THEN 2 ELSE 3 END, topic, question_id
        """,
        tuple(params),
    ).fetchall()


def parse_lines(text):
    if isinstance(text, list):
        return [str(item).strip() for item in text if str(item).strip()]
    text = "" if text is None else str(text)
    if "|" in text and "\n" not in text:
        return [item.strip() for item in text.split("|") if item.strip()]
    return [line.strip() for line in text.splitlines() if line.strip()]


def slugify(value):
    slug = re.sub(r"[^a-z0-9]+", "-", (value or "").strip().lower()).strip("-")
    return slug or "role"


def unique_role_slug(conn, desired_slug, current_role_id=None):
    base_slug = slugify(desired_slug)
    slug = base_slug
    counter = 2
    while True:
        params = [slug]
        extra_sql = ""
        if current_role_id:
            extra_sql = "AND role_id <> %s"
            params.append(current_role_id)
        row = conn.execute(
            f"SELECT 1 FROM roles WHERE role_slug = %s {extra_sql}",
            tuple(params),
        ).fetchone()
        if not row:
            return slug
        slug = f"{base_slug}-{counter}"
        counter += 1


def parse_bool(value, default=True):
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "yes", "true", "active", "enabled"}


def normalize_import_header(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def read_import_rows(uploaded_file):
    filename = secure_filename(uploaded_file.filename or "")
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        raw = uploaded_file.read().decode("utf-8-sig")
        return list(csv.DictReader(raw.splitlines()))
    if extension in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("Excel import requires openpyxl. Run: python -m pip install -r requirements.txt") from exc
        workbook = load_workbook(uploaded_file.stream, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [normalize_import_header(cell) for cell in rows[0]]
        return [dict(zip(headers, values)) for values in rows[1:] if any(value not in (None, "") for value in values)]
    raise ValueError("Please upload a .xlsx, .xlsm, or .csv file.")


def import_question_rows(conn, rows):
    imported = 0
    roles_created = 0
    role_slugs = set()

    for row_number, raw_row in enumerate(rows, start=2):
        row = {normalize_import_header(key): value for key, value in raw_row.items()}
        role_name = str(row.get("role_name") or row.get("role") or "").strip()
        role_slug = slugify(row.get("role_slug") or role_name)
        section = str(row.get("section") or "").strip()
        topic = str(row.get("topic") or "").strip()
        question_text = str(row.get("question_text") or row.get("question") or "").strip()
        correct_answer = str(row.get("correct_answer") or row.get("answer") or "").strip()

        if not role_name:
            role_name = role_slug.replace("-", " ").title()
        if not all([role_slug, section, topic, question_text, correct_answer]):
            raise ValueError(f"Row {row_number} is missing role, section, topic, question, or correct answer.")

        role_exists = conn.execute("SELECT 1 FROM roles WHERE role_slug = %s", (role_slug,)).fetchone()
        conn.execute(
            """
            INSERT INTO roles (role_slug, role_name)
            VALUES (%s, %s)
            ON CONFLICT (role_slug) DO UPDATE SET role_name = EXCLUDED.role_name
            """,
            (role_slug, role_name),
        )
        if not role_exists:
            roles_created += 1
        role_slugs.add(role_slug)

        question_code = str(row.get("question_code") or "").strip() or None
        difficulty = str(row.get("difficulty") or "Medium").strip()
        marks = int(row.get("marks") or 5)
        active = parse_bool(row.get("active"), default=True)
        options = parse_lines(row.get("options"))
        keywords = parse_lines(row.get("keywords"))

        if question_code:
            conn.execute(
                """
                INSERT INTO question_bank
                    (question_code, role_slug, section, topic, difficulty, question_text,
                     options, correct_answer, keywords, marks, active, updated_at)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, now())
                ON CONFLICT (question_code) DO UPDATE SET
                    role_slug = EXCLUDED.role_slug,
                    section = EXCLUDED.section,
                    topic = EXCLUDED.topic,
                    difficulty = EXCLUDED.difficulty,
                    question_text = EXCLUDED.question_text,
                    options = EXCLUDED.options,
                    correct_answer = EXCLUDED.correct_answer,
                    keywords = EXCLUDED.keywords,
                    marks = EXCLUDED.marks,
                    active = EXCLUDED.active,
                    updated_at = EXCLUDED.updated_at
                """,
                (
                    question_code,
                    role_slug,
                    section,
                    topic,
                    difficulty,
                    question_text,
                    Jsonb(options),
                    correct_answer,
                    Jsonb(keywords),
                    marks,
                    active,
                ),
            )
        else:
            conn.execute(
                """
                INSERT INTO question_bank
                    (question_code, role_slug, section, topic, difficulty, question_text,
                     options, correct_answer, keywords, marks, active, updated_at)
                VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, now())
                """,
                (
                    None,
                    role_slug,
                    section,
                    topic,
                    difficulty,
                    question_text,
                    Jsonb(options),
                    correct_answer,
                    Jsonb(keywords),
                    marks,
                    active,
                ),
            )
        imported += 1

    return {"questions": imported, "roles": roles_created, "role_slugs": sorted(role_slugs)}


@app.route("/", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        with get_db() as conn:
            ensure_hr_user_table(conn)
            user = conn.execute(
                """
                SELECT *
                FROM hr_portal_users
                WHERE username = %s AND active = TRUE
                """,
                (username,),
            ).fetchone()
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["hr_logged_in"] = True
            session["hr_username"] = user["username"]
            session["hr_full_name"] = user["full_name"]
            session["hr_is_admin"] = user["is_admin"]
            return redirect(request.args.get("next") or url_for("dashboard"))
        flash("Invalid HR credentials.", "error")
    return render_template("login.html", company_name=COMPANY_NAME)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/dashboard")
@require_login
def dashboard():
    stats, role_rows, recent = fetch_dashboard_stats()
    return render_template(
        "dashboard.html",
        stats=stats,
        role_rows=role_rows,
        recent=recent,
        company_name=COMPANY_NAME,
        db_label=mask_dsn_password(DATABASE_URL),
    )


@app.route("/candidates")
@require_login
def candidates():
    status = request.args.get("status", "")
    role_id = request.args.get("role_id", "")
    search = request.args.get("search", "").strip()
    shortlist_mode = request.args.get("shortlist_mode", "automatic")
    if shortlist_mode not in {"automatic", "manual"}:
        shortlist_mode = "automatic"
    with get_db() as conn:
        roles = fetch_roles(conn)
        reports = fetch_reports(conn, status=status, role_id=role_id, search=search)
    shortlisted_reports, not_shortlisted_reports = group_reports_by_shortlist(reports, mode=shortlist_mode)
    return render_template(
        "candidates.html",
        reports=reports,
        shortlisted_reports=shortlisted_reports,
        not_shortlisted_reports=not_shortlisted_reports,
        roles=roles,
        filters={"status": status, "role_id": role_id, "search": search, "shortlist_mode": shortlist_mode},
        shortlist_min_score=SHORTLIST_MIN_SCORE,
        shortlist_mode=shortlist_mode,
        company_name=COMPANY_NAME,
    )


@app.post("/candidates/<interview_id>/approve")
@require_login
def approve_candidate(interview_id):
    with get_db() as conn:
        candidate = conn.execute(
            """
            SELECT i.interview_id, i.total_score, u.name, u.email, u.phone, r.role_name
            FROM interviews i
            JOIN users u ON u.user_id = i.user_id
            JOIN roles r ON r.role_id = i.role_id
            WHERE i.interview_id = %s
            """,
            (interview_id,),
        ).fetchone()
        if not candidate:
            flash("Candidate interview was not found.", "error")
            return redirect(url_for("candidates", shortlist_mode="manual"))
    try:
        mail_result = send_candidate_shortlist_email(candidate)
        if not mail_result["sent"]:
            flash(f"{candidate['name']} was not approved because email could not be sent. {friendly_smtp_error(mail_result['error'])}", "error")
            return redirect(request.referrer or url_for("candidates", shortlist_mode="manual"))
    except Exception as exc:
        flash(f"{candidate['name']} was not approved because email could not be sent. {friendly_smtp_error(exc)}", "error")
        return redirect(request.referrer or url_for("candidates", shortlist_mode="manual"))

    with get_db() as conn:
        conn.execute(
            """
            UPDATE interviews
            SET shortlist_status = %s,
                shortlist_reason = %s
            WHERE interview_id = %s
            """,
            (
                "Shortlisted",
                f"Approved by HR for second round technical interview. Score: {candidate['total_score']}.",
                interview_id,
            ),
        )
    flash(f"{candidate['name']} approved and shortlist email sent.", "success")
    return redirect(request.referrer or url_for("candidates", shortlist_mode="manual"))


@app.route("/roles")
@require_login
def roles():
    with get_db() as conn:
        rows = fetch_role_question_counts(conn)
    return render_template(
        "roles.html",
        roles=rows,
        sections=SECTION_COUNTS,
        company_name=COMPANY_NAME,
    )


@app.route("/roles/new", methods=["GET", "POST"])
@require_login
def new_role():
    return save_role()


@app.route("/roles/<int:role_id>/edit", methods=["GET", "POST"])
@require_login
def edit_role(role_id):
    return save_role(role_id)


def save_role(role_id=None):
    with get_db() as conn:
        role = None
        if role_id:
            role = conn.execute("SELECT * FROM roles WHERE role_id = %s", (role_id,)).fetchone()
            if not role:
                flash("Role not found.", "error")
                return redirect(url_for("roles"))

        if request.method == "POST":
            role_name = request.form.get("role_name", "").strip()
            requested_slug = request.form.get("role_slug", "").strip() or role_name
            if not role_name:
                flash("Role name is required.", "error")
            else:
                role_slug = unique_role_slug(conn, requested_slug, current_role_id=role_id)
                if role_id:
                    old_slug = role["role_slug"]
                    conn.execute(
                        """
                        UPDATE roles
                        SET role_slug = %s, role_name = %s
                        WHERE role_id = %s
                        """,
                        (role_slug, role_name, role_id),
                    )
                    if old_slug != role_slug:
                        conn.execute(
                            "UPDATE question_bank SET role_slug = %s, updated_at = now() WHERE role_slug = %s",
                            (role_slug, old_slug),
                        )
                    flash("Role updated.", "success")
                else:
                    conn.execute(
                        "INSERT INTO roles (role_slug, role_name) VALUES (%s, %s)",
                        (role_slug, role_name),
                    )
                    flash("Role created. You can now add questions for it.", "success")
                return redirect(url_for("roles"))

    return render_template(
        "role_form.html",
        role=role,
        company_name=COMPANY_NAME,
    )


@app.route("/users")
@require_admin
def users():
    with get_db() as conn:
        ensure_hr_user_table(conn)
        rows = conn.execute(
            """
            SELECT user_id, username, email, full_name, is_admin, active, must_change_password, created_at
            FROM hr_portal_users
            ORDER BY is_admin DESC, full_name
            """
        ).fetchall()
    return render_template(
        "users.html",
        users=rows,
        company_name=COMPANY_NAME,
    )


@app.route("/users/new", methods=["GET", "POST"])
@require_admin
def new_user():
    created_password = ""
    mail_result = None
    form_values = {"full_name": "", "email": "", "username": ""}
    if request.method == "POST":
        full_name = request.form.get("full_name", "").strip()
        email = request.form.get("email", "").strip().lower()
        username = request.form.get("username", "").strip() or email
        password = request.form.get("password", "")
        confirm_password = request.form.get("confirm_password", "")
        is_admin = bool(request.form.get("is_admin"))
        form_values = {"full_name": full_name, "email": email, "username": username}

        if not all([full_name, email, username, password, confirm_password]):
            flash("Name, email, username, password, and confirm password are required.", "error")
        elif password != confirm_password:
            flash("Password and confirm password do not match.", "error")
        elif len(password) < 8:
            flash("Password must be at least 8 characters.", "error")
        else:
            created_password = password
            try:
                with get_db() as conn:
                    ensure_hr_user_table(conn)
                    user = conn.execute(
                        """
                        INSERT INTO hr_portal_users
                            (username, email, full_name, password_hash, is_admin, active, must_change_password)
                        VALUES (%s, %s, %s, %s, %s, TRUE, FALSE)
                        RETURNING username, email, full_name
                        """,
                        (username, email, full_name, generate_password_hash(password), is_admin),
                    ).fetchone()
                try:
                    mail_result = send_hr_access_email(user, password)
                except Exception as exc:
                    mail_result = {"sent": False, "status": "failed", "error": str(exc)}
                if mail_result["sent"]:
                    flash("HR user created and username/password email sent.", "success")
                    return redirect(url_for("users"))
                flash("HR user created. Email was not sent, so share the username and password manually.", "success")
            except psycopg.errors.UniqueViolation:
                flash("A user with that username or email already exists.", "error")
                created_password = ""

    return render_template(
        "user_form.html",
        created_password=created_password,
        form_values=form_values,
        mail_result=mail_result,
        portal_url=HR_PORTAL_URL,
        company_name=COMPANY_NAME,
    )


@app.post("/users/<int:user_id>/toggle")
@require_admin
def toggle_user(user_id):
    with get_db() as conn:
        ensure_hr_user_table(conn)
        row = conn.execute("SELECT username, active FROM hr_portal_users WHERE user_id = %s", (user_id,)).fetchone()
        if not row:
            flash("User not found.", "error")
        elif row["username"] == session.get("hr_username"):
            flash("You cannot disable your own account.", "error")
        else:
            conn.execute(
                "UPDATE hr_portal_users SET active = %s, updated_at = now() WHERE user_id = %s",
                (not row["active"], user_id),
            )
            flash("User access updated.", "success")
    return redirect(url_for("users"))


@app.post("/users/<int:user_id>/reset-password")
@require_admin
def reset_user_password(user_id):
    temporary_password = secrets.token_urlsafe(9)
    with get_db() as conn:
        ensure_hr_user_table(conn)
        user = conn.execute(
            "SELECT username, email, full_name FROM hr_portal_users WHERE user_id = %s",
            (user_id,),
        ).fetchone()
        if not user:
            flash("User not found.", "error")
            return redirect(url_for("users"))
        conn.execute(
            """
            UPDATE hr_portal_users
            SET password_hash = %s, must_change_password = TRUE, updated_at = now()
            WHERE user_id = %s
            """,
            (generate_password_hash(temporary_password), user_id),
        )
    try:
        mail_result = send_hr_access_email(user, temporary_password)
        if mail_result["sent"]:
            flash("Temporary password reset and emailed.", "success")
        else:
            flash(f"Temporary password reset: {temporary_password}", "success")
    except Exception as exc:
        flash(f"Temporary password reset: {temporary_password}. Email failed: {exc}", "error")
    return redirect(url_for("users"))


@app.route("/settings/smtp", methods=["GET", "POST"])
@require_admin
def smtp_settings():
    settings = {
        "SMTP_HOST": SMTP_HOST,
        "SMTP_PORT": str(SMTP_PORT),
        "SMTP_USER": SMTP_USER,
        "SMTP_PASSWORD": "",
        "SMTP_FROM": SMTP_FROM,
    }
    if request.method == "POST":
        submitted = {
            "SMTP_HOST": request.form.get("smtp_host", "").strip(),
            "SMTP_PORT": request.form.get("smtp_port", "587").strip() or "587",
            "SMTP_USER": request.form.get("smtp_user", "").strip(),
            "SMTP_FROM": request.form.get("smtp_from", "").strip(),
        }
        password = request.form.get("smtp_password", "")
        if password:
            submitted["SMTP_PASSWORD"] = password.replace(" ", "").strip()
        else:
            submitted["SMTP_PASSWORD"] = SMTP_PASSWORD

        if not all([submitted["SMTP_HOST"], submitted["SMTP_PORT"], submitted["SMTP_USER"], submitted["SMTP_PASSWORD"], submitted["SMTP_FROM"] or submitted["SMTP_USER"]]):
            flash("SMTP host, port, user, password, and from email are required.", "error")
        else:
            try:
                int(submitted["SMTP_PORT"])
                update_env_file(submitted)
                apply_smtp_settings(submitted)
                flash("SMTP settings saved.", "success")
                if request.form.get("send_test") == "yes":
                    test_user = {
                        "email": SMTP_FROM,
                        "full_name": session.get("hr_full_name") or session.get("hr_username") or "HR Admin",
                        "username": session.get("hr_username") or "hr",
                    }
                    result = send_hr_access_email(test_user, "SMTP test only")
                    if result["sent"]:
                        flash("SMTP test email sent successfully.", "success")
                    else:
                        flash(f"SMTP test email was not sent. {friendly_smtp_error(result['error'])}", "error")
                return redirect(url_for("smtp_settings"))
            except Exception as exc:
                flash(f"SMTP settings saved, but test email failed. {friendly_smtp_error(exc)}", "error")

        settings.update(
            {
                "SMTP_HOST": submitted["SMTP_HOST"],
                "SMTP_PORT": submitted["SMTP_PORT"],
                "SMTP_USER": submitted["SMTP_USER"],
                "SMTP_FROM": submitted["SMTP_FROM"],
            }
        )

    configured = all([SMTP_HOST, SMTP_USER, SMTP_PASSWORD, SMTP_FROM])
    return render_template(
        "smtp_settings.html",
        settings=settings,
        configured=configured,
        company_name=COMPANY_NAME,
    )


@app.route("/questions")
@require_login
def questions():
    role_slug = request.args.get("role_slug", "")
    section = request.args.get("section", "")
    active = request.args.get("active", "")
    with get_db() as conn:
        roles = fetch_roles(conn)
        rows = fetch_questions(conn, role_slug=role_slug, section=section, active=active)
    return render_template(
        "questions.html",
        questions=rows,
        roles=roles,
        sections=SECTION_COUNTS,
        filters={"role_slug": role_slug, "section": section, "active": active},
        company_name=COMPANY_NAME,
    )


@app.route("/questions/new", methods=["GET", "POST"])
@require_login
def new_question():
    return save_question()


@app.route("/questions/import", methods=["GET", "POST"])
@require_login
def import_questions():
    if request.method == "POST":
        uploaded_file = request.files.get("question_file")
        if not uploaded_file or not uploaded_file.filename:
            flash("Please choose an Excel or CSV file.", "error")
        else:
            try:
                rows = read_import_rows(uploaded_file)
                if not rows:
                    raise ValueError("The uploaded file does not contain any question rows.")
                with get_db() as conn:
                    result = import_question_rows(conn, rows)
                flash(
                    f"Imported {result['questions']} questions across {len(result['role_slugs'])} role(s). "

                    f"New roles created: {result['roles']}.",
                    "success",
                )
                return redirect(url_for("questions"))
            except Exception as exc:
                flash(str(exc), "error")
    return render_template(
        "question_import.html",
        columns=QUESTION_IMPORT_COLUMNS,
        company_name=COMPANY_NAME,
    )


@app.route("/questions/<int:question_id>/edit", methods=["GET", "POST"])
@require_login
def edit_question(question_id):
    return save_question(question_id)


def save_question(question_id=None):
    with get_db() as conn:
        roles = fetch_roles(conn)
        question = None
        selected_role_slug = request.args.get("role_slug", "").strip()
        if question_id:
            question = conn.execute("SELECT * FROM question_bank WHERE question_id = %s", (question_id,)).fetchone()
            if not question:
                flash("Question not found.", "error")
                return redirect(url_for("questions"))
            selected_role_slug = question["role_slug"]

        if request.method == "POST":
            role_slug = request.form.get("role_slug", "").strip()
            section = request.form.get("section", "").strip()
            topic = request.form.get("topic", "").strip()
            question_code = request.form.get("question_code", "").strip() or None
            difficulty = request.form.get("difficulty", "Medium").strip()
            question_text = request.form.get("question_text", "").strip()
            correct_answer = request.form.get("correct_answer", "").strip()
            options = parse_lines(request.form.get("options", ""))
            keywords = parse_lines(request.form.get("keywords", ""))
            marks = int(request.form.get("marks", 5) or 5)
            active = bool(request.form.get("active"))

            if not all([role_slug, section, topic, question_text, correct_answer]):
                flash("Role, section, topic, question, and answer are required.", "error")
            elif question_id:
                conn.execute(
                    """
                    UPDATE question_bank
                    SET question_code = %s, role_slug = %s, section = %s, topic = %s,
                        difficulty = %s, question_text = %s, options = %s, correct_answer = %s,
                        keywords = %s, marks = %s, active = %s, updated_at = now()
                    WHERE question_id = %s
                    """,
                    (
                        question_code,
                        role_slug,
                        section,
                        topic,
                        difficulty,
                        question_text,
                        Jsonb(options),
                        correct_answer,
                        Jsonb(keywords),
                        marks,
                        active,
                        question_id,
                    ),
                )
                flash("Question updated.", "success")
                return redirect(url_for("questions", role_slug=role_slug, section=section))
            else:
                conn.execute(
                    """
                    INSERT INTO question_bank
                        (question_code, role_slug, section, topic, difficulty, question_text,
                         options, correct_answer, keywords, marks, active, updated_at)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, now())
                    """,
                    (
                        question_code,
                        role_slug,
                        section,
                        topic,
                        difficulty,
                        question_text,
                        Jsonb(options),
                        correct_answer,
                        Jsonb(keywords),
                        marks,
                        active,
                    ),
                )
                flash("Question created.", "success")
                return redirect(url_for("questions", role_slug=role_slug, section=section))

    return render_template(
        "question_form.html",
        question=question,
        roles=roles,
        selected_role_slug=selected_role_slug,
        sections=SECTION_COUNTS,
        company_name=COMPANY_NAME,
    )


@app.post("/questions/<int:question_id>/toggle")
@require_login
def toggle_question(question_id):
    with get_db() as conn:
        row = conn.execute("SELECT active FROM question_bank WHERE question_id = %s", (question_id,)).fetchone()
        if row:
            conn.execute("UPDATE question_bank SET active = %s, updated_at = now() WHERE question_id = %s", (not row["active"], question_id))
            flash("Question status changed.", "success")
    return redirect(request.referrer or url_for("questions"))


@app.route("/download/report/<interview_id>")
@require_login
def download_report(interview_id):
    with get_db() as conn:
        row = conn.execute("SELECT report_path FROM interviews WHERE interview_id = %s", (interview_id,)).fetchone()
    path = safe_project_file(row["report_path"] if row else "", REPORT_DIR)
    if not path:
        flash("Report file is not available.", "error")
        return redirect(url_for("candidates"))
    return send_file(path, as_attachment=True)


@app.route("/download/resume/<interview_id>")
@require_login
def download_resume(interview_id):
    with get_db() as conn:
        row = conn.execute(
            """
            SELECT u.resume_path
            FROM interviews i
            JOIN users u ON u.user_id = i.user_id
            WHERE i.interview_id = %s
            """,
            (interview_id,),
        ).fetchone()
    path = safe_project_file(row["resume_path"] if row else "", RESUME_DIR)
    if not path:
        flash("Resume file is not available.", "error")
        return redirect(url_for("candidates"))
    return send_file(path, as_attachment=True)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "5050")), debug=True)
